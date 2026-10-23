"""DOCX parsing with heading styles and tables preserved.

DOCX has no fixed pages; we emit a single logical page per ~3000 chars so
citations stay stable, and mark heading lines with markdown-style '#'.
"""
from __future__ import annotations

from pathlib import Path

from backend.app.ingestion.hashes import sha256_file
from backend.app.schemas.documents import ParsedDocument, ParsedPage, TableBlock

PAGE_CHARS = 3000


def _iter_block_items(doc):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def parse_docx(path: str | Path, doc_id: str) -> ParsedDocument:
    import docx  # python-docx
    from docx.table import Table

    path = Path(path)
    document = docx.Document(str(path))
    lines: list[str] = []
    tables: list[dict] = []
    for block in _iter_block_items(document):
        if isinstance(block, Table):
            rows = [[c.text.strip() for c in r.cells] for r in block.rows]
            if rows:
                tables.append(TableBlock(headers=rows[0], rows=rows[1:]).model_dump())
                lines.append("")
                for r in rows:
                    lines.append(" | ".join(r))
                lines.append("")
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        if style.lower().startswith("heading"):
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            lines.append(f"{'#' * min(level, 6)} {text}")
        elif style.lower() == "title":
            lines.append(f"# {text}")
        else:
            lines.append(text)

    pages: list[ParsedPage] = []
    buf: list[str] = []
    size = 0
    for line in lines:
        buf.append(line)
        size += len(line) + 1
        if size >= PAGE_CHARS and not line.startswith("#"):
            pages.append(ParsedPage(page=len(pages) + 1, text="\n".join(buf).strip()))
            buf, size = [], 0
    if buf or not pages:
        pages.append(ParsedPage(page=len(pages) + 1, text="\n".join(buf).strip()))
    props = document.core_properties
    meta = {"title": props.title, "author": props.author, "tables": tables,
            "page_count": len(pages), "logical_pages": True}
    return ParsedDocument(doc_id=doc_id, filename=path.name, file_sha256=sha256_file(path),
                          parser="docx", pages=pages, metadata=meta)


def parse_text_file(path: str | Path, doc_id: str) -> ParsedDocument:
    path = Path(path)
    text = path.read_text(encoding="utf-8", errors="replace")
    pages: list[ParsedPage] = []
    for i in range(0, max(len(text), 1), PAGE_CHARS * 2):
        pages.append(ParsedPage(page=len(pages) + 1, text=text[i:i + PAGE_CHARS * 2]))
    return ParsedDocument(doc_id=doc_id, filename=path.name, file_sha256=sha256_file(path),
                          parser="text", pages=pages, metadata={"page_count": len(pages),
                                                                "logical_pages": True})
