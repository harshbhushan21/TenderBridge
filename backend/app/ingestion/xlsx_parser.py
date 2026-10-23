"""XLSX parsing: every sheet becomes a page; BoQ rows become structured items.

Cell coordinates are preserved so commercial tools can cite `Sheet!B12`.
No arithmetic is done here beyond reading cached values.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from backend.app.ingestion.hashes import sha256_file
from backend.app.schemas.documents import BoqItem, ParsedDocument, ParsedPage, TableBlock

_HEADER_ALIASES = {
    "item_no": ("item no", "item", "s.no", "s. no", "sl no", "sl. no", "sr no", "sr. no", "item code", "#", "no"),
    "description": ("description", "item description", "particulars", "service", "details", "scope"),
    "unit": ("unit", "uom", "units"),
    "quantity": ("qty", "quantity", "qnty", "nos"),
    "rate": ("rate", "unit rate", "unit price", "price", "rate (inr)", "unit cost"),
    "amount": ("amount", "total", "total amount", "value", "amount (inr)", "total price"),
}


def _norm(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip().lower()


def _to_float(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[₹,\s]|Rs\.?|INR", "", str(v), flags=re.I)
    try:
        return float(s)
    except ValueError:
        return None


def _map_headers(row: list[Any]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, h in enumerate(row):
        n = _norm(h)
        if not n:
            continue
        for field, aliases in _HEADER_ALIASES.items():
            if field in mapping:
                continue
            if n in aliases or any(n.startswith(a) for a in aliases if len(a) > 3):
                mapping[field] = idx
                break
    return mapping


def parse_xlsx(path: str | Path, doc_id: str) -> tuple[ParsedDocument, list[BoqItem]]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    path = Path(path)
    wb = load_workbook(str(path), data_only=True, read_only=True)
    pages: list[ParsedPage] = []
    items: list[BoqItem] = []
    tables: list[dict] = []
    for sheet in wb.worksheets:
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        lines = [f"# Sheet: {sheet.title}"]
        header_idx, mapping = None, {}
        for i, r in enumerate(rows[:25]):
            m = _map_headers(r)
            if "description" in m and len(m) >= 2:
                header_idx, mapping = i, m
                break
        for i, r in enumerate(rows):
            if all(v is None or str(v).strip() == "" for v in r):
                continue
            lines.append(" | ".join("" if v is None else str(v) for v in r))
        if header_idx is not None:
            headers = ["" if v is None else str(v) for v in rows[header_idx]]
            body: list[list[str]] = []
            for ri in range(header_idx + 1, len(rows)):
                r = rows[ri]
                if all(v is None or str(v).strip() == "" for v in r):
                    continue
                desc_v = r[mapping["description"]] if mapping["description"] < len(r) else None
                if desc_v is None or str(desc_v).strip() == "":
                    continue
                excel_row = ri + 1
                cells = {f"{get_column_letter(ci + 1)}{excel_row}": v for ci, v in enumerate(r)
                         if v is not None}

                def g(field: str):
                    idx = mapping.get(field)
                    return r[idx] if idx is not None and idx < len(r) else None

                items.append(BoqItem(
                    sheet=sheet.title, row=excel_row,
                    item_no=None if g("item_no") is None else str(g("item_no")),
                    description=str(desc_v).strip(), unit=None if g("unit") is None else str(g("unit")),
                    quantity=_to_float(g("quantity")), rate=_to_float(g("rate")),
                    amount=_to_float(g("amount")), cells={k: (v if isinstance(v, (int, float)) else str(v))
                                                          for k, v in cells.items()}))
                body.append(["" if v is None else str(v) for v in r])
            tables.append(TableBlock(caption=sheet.title, headers=headers, rows=body,
                                     page=len(pages) + 1).model_dump())
        pages.append(ParsedPage(page=len(pages) + 1, text="\n".join(lines)))
    wb.close()
    doc = ParsedDocument(doc_id=doc_id, filename=path.name, file_sha256=sha256_file(path),
                         parser="xlsx", pages=pages,
                         metadata={"sheets": [p.text.splitlines()[0][9:] for p in pages],
                                   "tables": tables, "page_count": len(pages), "logical_pages": True})
    return doc, items
