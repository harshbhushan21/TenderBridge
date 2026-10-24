"""Compliance matrix and eligibility/risk workbooks (openpyxl)."""
from __future__ import annotations

from collections.abc import Iterable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.app.schemas.assessment import BidAssessment
from backend.app.schemas.documents import SourceSpan
from backend.app.schemas.eligibility import ComplianceItem, RiskItem
from backend.app.schemas.requirements import KeyFacts, Requirement, RequirementCategory

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
WARN_FONT = Font(bold=True, color="9C0006")

STATUS_FILLS = {
    "meets": "C6EFCE",
    "meets_with_undertaking": "E2EFDA",
    "meets_with_substitution": "E2EFDA",
    "needs_action": "FFEB9C",
    "partial": "FFEB9C",
    "needs_human": "FCE4D6",
    "gap": "FFC7CE",
    "within_policy": "C6EFCE",
    "review": "FFEB9C",
    "escalate": "FFC7CE",
    "not_analysed": "EDEDED",
}

SUBMISSION_CATEGORIES = {RequirementCategory.submission.value, RequirementCategory.eligibility_registration.value}


def _val(x: Any) -> str:
    return getattr(x, "value", x) if x is not None else ""


def _span_text(span: SourceSpan | None, max_quote: int = 200) -> str:
    if span is None:
        return ""
    quote = (span.text_span or "").replace("\n", " ").strip()
    if len(quote) > max_quote:
        quote = quote[: max_quote - 1] + "…"
    return f"{span.short()}: \"{quote}\"" if quote else span.short()


def _spans_text(spans: Iterable[SourceSpan]) -> str:
    return "\n".join(_span_text(s, 120) for s in spans)


def _header_block(ws: Worksheet, title: str, *, tender_id: str, version: int,
                  analysis_mode: str = "focused") -> int:
    """Write title rows; return next free row."""
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Tender: {tender_id}    Version: v{version}    "
                                   f"Generated: {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}")
    mode_txt = ("Focused live analysis — not full-tender coverage" if analysis_mode == "focused"
                else "Full-tender analysis")
    ws.cell(row=3, column=1, value=f"Analysis mode: {mode_txt}")
    ws.cell(row=4, column=1, value="SYNTHETIC DEMO DATA — not a real tender or bidder. "
                                   "Recommendation only; human approval required. No portal submission.").font = WARN_FONT
    return 6


_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _table(ws: Worksheet, start_row: int, headers: list[str], rows: list[list[Any]],
           *, status_col: int | None = None, widths: list[int] | None = None) -> None:
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=ci, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for ri, row in enumerate(rows, start=start_row + 1):
        for ci, v in enumerate(row, start=1):
            c = ws.cell(row=ri, column=ci, value=v)
            if isinstance(v, str) and v[:1] in _FORMULA_START:
                c.data_type = "s"  # tender text is data: never let "=HYPERLINK(...)" become a formula
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if status_col is not None:
            st = str(row[status_col - 1] or "")
            fill = STATUS_FILLS.get(st)
            if fill:
                ws.cell(row=ri, column=status_col).fill = PatternFill("solid", fgColor=fill)
    last_col = get_column_letter(len(headers))
    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)  # type: ignore[assignment]
    ws.auto_filter.ref = f"A{start_row}:{last_col}{max(start_row, start_row + len(rows))}"
    for ci, h in enumerate(headers, start=1):
        w = (widths[ci - 1] if widths and ci - 1 < len(widths) else max(12, min(40, len(h) + 4)))
        ws.column_dimensions[get_column_letter(ci)].width = w


def write_compliance_xlsx(path: Path | str, *, tender_id: str, version: int,
                          requirements: list[Requirement], compliance: list[ComplianceItem],
                          risks: list[RiskItem], key_facts: KeyFacts | None = None,
                          analysis_mode: str = "focused") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    by_req = {c.requirement_id: c for c in compliance}
    wb = Workbook()

    ws = cast(Worksheet, wb.active)
    ws.title = "Compliance Matrix"
    start = _header_block(ws, "Compliance Matrix", tender_id=tender_id, version=version,
                          analysis_mode=analysis_mode)
    headers = ["Requirement ID", "Clause", "Page", "Mandatory", "Hard gate", "Category", "Status",
               "Summary", "Tender evidence", "Company evidence", "Owner", "Actions", "Version"]
    rows: list[list[Any]] = []
    for r in requirements:
        c = by_req.get(r.requirement_id)
        rows.append([
            r.requirement_id, r.clause_ref or "", r.source.page, "Yes" if r.mandatory else "No",
            "Yes" if (c.hard_gate if c else r.hard_gate) else "No", _val(r.category),
            _val(c.status) if c else "not_analysed",
            (c.summary if c else r.title or r.text[:200]),
            _span_text(c.tender_source if c else r.source),
            _spans_text(c.company_evidence) if c else "",
            (c.owner or "") if c else "",
            "\n".join(c.actions) if c else "",
            f"v{r.tender_version}" + (" (amended)" if r.amended else ""),
        ])
    _table(ws, start, headers, rows, status_col=7,
           widths=[16, 10, 7, 10, 9, 22, 18, 50, 50, 45, 14, 40, 10])

    ck = cast(Worksheet, wb.create_sheet("Submission Checklist"))
    start = _header_block(ck, "Submission Checklist", tender_id=tender_id, version=version,
                          analysis_mode=analysis_mode)
    ck_rows: list[list[Any]] = []
    for r in requirements:
        if _val(r.category) not in SUBMISSION_CATEGORIES:
            continue
        c = by_req.get(r.requirement_id)
        ck_rows.append([r.requirement_id, r.clause_ref or "", r.title or r.text[:120], r.source.short(),  # type: ignore[union-attr]
                        _val(c.status) if c else "not_analysed", "\n".join(c.actions) if c else "", "☐"])
    if key_facts and key_facts.submission_deadline:
        deadline_source = key_facts.source_for("submission_deadline")
        ck_rows.insert(0, ["KEY", "", f"Submit before {key_facts.submission_deadline.isoformat()}",
                           deadline_source.short() if deadline_source else "", "", "", "☐"])
    _table(ck, start, ["Requirement ID", "Clause", "Item", "Source", "Status", "Actions", "Done"],
           ck_rows, status_col=5, widths=[16, 10, 60, 22, 18, 40, 8])

    if risks:
        rk = cast(Worksheet, wb.create_sheet("Commercial Risks"))
        start = _header_block(rk, "Commercial Risks", tender_id=tender_id, version=version,
                              analysis_mode=analysis_mode)
        _table(rk, start, *_risk_rows(risks), status_col=3, widths=[14, 20, 14, 30, 30, 50, 12, 14, 30])

    wb.save(path)
    return path


def _risk_rows(risks: list[RiskItem]) -> tuple[list[str], list[list[Any]]]:
    headers = ["Risk ID", "Category", "Level", "Tender value", "Policy value", "Summary", "Approver",
               "Approval", "Source"]
    rows = [[r.risk_id, r.category, _val(r.level), r.tender_value, r.policy_value, r.summary,
             r.approver or "", r.approval_status, _span_text(r.source, 100)] for r in risks]
    return headers, rows


def write_eligibility_xlsx(path: Path | str, *, tender_id: str, version: int,
                           compliance: list[ComplianceItem], risks: list[RiskItem],
                           assessment: BidAssessment | None,
                           analysis_mode: str | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    mode = analysis_mode or (assessment.analysis_mode if assessment else "focused")
    wb = Workbook()

    ws = cast(Worksheet, wb.active)
    ws.title = "Eligibility"
    start = _header_block(ws, "Eligibility Assessment", tender_id=tender_id, version=version,
                          analysis_mode=mode)
    rows: list[list[Any]] = []
    for c in compliance:
        if not c.category.startswith("eligibility") and c.category != "personnel" and not c.hard_gate:
            continue
        rr = c.rule_result
        actual = "\n".join(f"{k.label}: {k.value}{(' ' + k.unit) if k.unit else ''}"
                           for k in (rr.calculations if rr else []))
        rows.append([c.requirement_id, rr.rule_kind if rr else c.category, "Yes" if c.hard_gate else "No",
                     actual, _val(c.status), c.summary,
                     _spans_text([c.tender_source] + list(c.company_evidence)),
                     "\n".join(c.actions or (rr.actions if rr else [])),
                     (rr.needs_human_reason or "") if rr else "",
                     f"v{c.tender_version}" + (f" (reused v{c.reused_from_version})" if c.reused_from_version else "")])
    _table(ws, start, ["Requirement ID", "Rule", "Hard gate", "Actual company value", "Verdict", "Summary",
                       "Proof", "Unresolved actions", "Needs human", "Version"],
           rows, status_col=5, widths=[16, 16, 9, 40, 18, 50, 55, 40, 30, 16])

    rk = cast(Worksheet, wb.create_sheet("Risk Register"))
    start = _header_block(rk, "Risk Register", tender_id=tender_id, version=version, analysis_mode=mode)
    _table(rk, start, *_risk_rows(risks), status_col=3, widths=[14, 20, 14, 30, 30, 50, 12, 14, 30])

    bs = cast(Worksheet, wb.create_sheet("Bid Score"))
    start = _header_block(bs, "Bid / No-Bid Score", tender_id=tender_id, version=version, analysis_mode=mode)
    if assessment:
        summary: list[list[Any]] = [
            ["Recommendation", _val(assessment.recommendation)],
            ["Weighted score", assessment.score],
            ["Hard gate failures", "\n".join(assessment.hard_gate_failures) or "None"],
            ["Conditions", "\n".join(assessment.conditions) or "None"],
            ["Policy version", assessment.policy_version],
            ["Run", assessment.analysis_run_id],
            ["Note", "Mandatory gates override the weighted score."],
        ]
        for i, (k, v) in enumerate(summary):
            bs.cell(row=start + i, column=1, value=k).font = Font(bold=True)
            bs.cell(row=start + i, column=2, value=v).alignment = Alignment(wrap_text=True)
        crow = start + len(summary) + 1
        _table(bs, crow, ["Criterion", "Label", "Weight", "Raw score", "Weighted", "Rationale"],
               [[c.key, c.label, c.weight, c.raw_score, c.weighted, c.rationale] for c in assessment.criteria],
               widths=[22, 30, 10, 12, 12, 60])
        grow = crow + len(assessment.criteria) + 3
        _table(bs, grow, ["Hard gate", "Label", "Status", "Passed", "Summary", "Evidence"],
               [[g.requirement_id, g.label, g.status, "Yes" if g.passed else "NO", g.summary,
                 _spans_text(g.evidence)] for g in assessment.hard_gates],
               status_col=3, widths=[22, 30, 18, 12, 50, 60])
        bs.freeze_panes = None
    else:
        bs.cell(row=start, column=1, value="No assessment available for this version.")

    wb.save(path)
    return path
