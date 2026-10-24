"""Phase 9: every export opens in its real reader and shows the tender version and readable citations.

v1 is approved at Gate 1, Corrigendum No. 1 makes v2, and v2 is decided NO_BID; exports are
downloaded through the API as a user would.
"""
from __future__ import annotations

import io
import json
import time
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.db.engine import session_scope
from backend.app.exports.citations import doc_labels
from backend.app.exports.draft_docx import write_draft_docx
from backend.app.ingestion.kb_ingest import ingest_company_kb
from backend.app.knowledge.kb_seed import seed_company_facts
from backend.app.main import create_app
from backend.app.schemas.assessment import ApprovalStatus, Claim, SectionDraft
from backend.app.schemas.documents import SourceSpan

DATA = Path(__file__).resolve().parents[3] / "demo_data"
ORIGINAL = DATA / "tender/TDR-001/original"
MANAGER = {"X-User": "bid.manager@nimbus.example", "X-Role": "bid_manager"}
TENDER_FILES = ("notice.pdf", "main_rfp.pdf", "annexures.pdf", "price_schedule_boq.xlsx")


def _wait(c: TestClient, tid: str, run_id: str, statuses: set[str]) -> str:
    end = time.monotonic() + 180
    while time.monotonic() < end:
        status = c.get(f"/tenders/{tid}/runs/{run_id}").json()["status"]
        if status in statuses:
            return status
        time.sleep(0.25)
    raise TimeoutError(run_id)


@pytest.fixture(scope="module")
def tender():
    with session_scope() as s:
        company = seed_company_facts(s, facts_dir=DATA / "structured_facts")["company_id"]
        ingest_company_kb(s, company, kb_dir=DATA / "company_kb", facts_dir=DATA / "structured_facts",
                          manifest_path=DATA / "demo_config/company_manifest.json")
    with TestClient(create_app()) as c:
        tid = c.post("/tenders", json={"title": "Exports test"}, headers=MANAGER).json()["tenderId"]
        c.post(f"/tenders/{tid}/files", headers=MANAGER,
               files=[("files", (n, (ORIGINAL / n).read_bytes())) for n in TENDER_FILES])
        run1 = c.post(f"/tenders/{tid}/run", json={}, headers=MANAGER).json()["analysis_run_id"]
        _wait(c, tid, run1, {"awaiting_human"})
        c.post(f"/tenders/{tid}/decision", json={"tender_version": 1, "action": "approve_bid"}, headers=MANAGER)
        _wait(c, tid, run1, {"completed"})
        aid = c.post(f"/tenders/{tid}/amendments", headers=MANAGER, files={
            "file": ("corrigendum.pdf", (ORIGINAL / "corrigendum.pdf").read_bytes())}).json()["amendment_id"]
        c.post(f"/tenders/{tid}/amendments/{aid}/confirm", json={}, headers=MANAGER)
        run2 = c.post(f"/tenders/{tid}/reanalyse", json={"mode": "impact"}, headers=MANAGER).json()["analysis_run_id"]
        _wait(c, tid, run2, {"awaiting_human"})
        c.post(f"/tenders/{tid}/decision", headers=MANAGER,
               json={"tender_version": 2, "action": "no_bid", "comment": "ISO 27001 expires before the deadline"})
        _wait(c, tid, run2, {"completed"})
        yield c, tid


def _download(c: TestClient, tid: str, kind: str, version: int) -> bytes:
    r = c.get(f"/tenders/{tid}/exports/{kind}", params={"version": version}, headers=MANAGER)
    assert r.status_code == 200, (kind, version, r.text[:200])
    return r.content


def _sheet_text(content: bytes) -> dict[str, list[list[str]]]:
    wb = load_workbook(io.BytesIO(content))
    return {ws.title: [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
            for ws in wb.worksheets}


@pytest.mark.parametrize("version", [1, 2])
def test_compliance_matrix_versioned_and_cited(tender, version):
    c, tid = tender
    sheets = _sheet_text(_download(c, tid, "compliance_xlsx", version))
    assert set(sheets) == {"Compliance Matrix", "Submission Checklist", "Commercial Risks"}
    for rows in sheets.values():
        assert f"Version: v{version}" in rows[1][0] and "SYNTHETIC DEMO DATA" in rows[3][0]
    matrix = sheets["Compliance Matrix"]
    header = next(i for i, r in enumerate(matrix) if r[0] == "Requirement ID")
    col = matrix[header].index("Tender evidence")
    cited = [r[col] for r in matrix[header + 1:] if r[col]]
    assert len(cited) == len(matrix) - header - 1  # every requirement carries its tender source
    assert any(x.startswith("main_rfp.pdf p.") for x in cited)
    assert not any(x.startswith("DOC-") for x in cited)  # filenames, not opaque ids
    if version == 2:
        checklist = sheets["Submission Checklist"]
        assert any("2026-11-20T15:00" in r[2] and r[3].startswith("corrigendum.pdf") for r in checklist)


@pytest.mark.parametrize("version", [1, 2])
def test_eligibility_workbook_versioned_with_proof(tender, version):
    c, tid = tender
    sheets = _sheet_text(_download(c, tid, "eligibility_xlsx", version))
    rows = sheets["Eligibility"]
    header = next(i for i, r in enumerate(rows) if r[0] == "Requirement ID")
    cols = {name: rows[header].index(name) for name in ("Hard gate", "Verdict", "Proof", "Version")}
    gates = [r for r in rows[header + 1:] if r[cols["Hard gate"]] == "Yes"]
    assert len(gates) == 13 and all(r[cols["Version"]].startswith(f"v{version}") for r in gates)
    assert all(r[cols["Proof"]] for r in gates)
    pq06 = next(r for r in gates if r[0] == "REQ-PQ-06")
    assert pq06[cols["Verdict"]] == ("meets" if version == 1 else "gap")
    assert {"Risk Register", "Bid Score"} <= set(sheets)


def test_calendar_uses_effective_dates_and_cites_the_corrigendum(tender):
    from ics import Calendar

    c, tid = tender
    for version, deadline in ((1, "2026-11-06"), (2, "2026-11-20")):
        cal = Calendar(_download(c, tid, "dates_ics", version).decode())
        events = {e.name: e for e in cal.events}
        sub = next(e for name, e in events.items() if name.endswith("Bid submission deadline"))
        assert sub.name.startswith(f"{tid} v{version}")
        assert sub.begin.to("Asia/Kolkata").date().isoformat() == deadline
        assert f"version v{version}" in sub.description
        assert ("corrigendum.pdf p.1" if version == 2 else "Source: ") in sub.description


def test_audit_json_records_both_decisions(tender):
    c, tid = tender
    data = json.loads(_download(c, tid, "audit_json", 2))
    assert data["tender_id"] == tid and data["tender_version"] == 2 and data["synthetic_demo_data"]
    actions = [e["action"] for e in data["audit_log"]]
    assert {"gate1.approve_bid", "amendment.confirmed", "gate1.no_bid"} <= set(actions)
    assert {(r["version"], r["action"], r["status"]) for r in data["reviews"]} >= {
        (1, "approve_bid", "superseded"), (2, "no_bid", "rejected")}
    assert all(item["tender_source"]["doc_id"] for item in data["compliance"])


def test_no_bid_report_only_for_no_bid_and_cites_failed_gate(tender):
    c, tid = tender
    assert c.get(f"/tenders/{tid}/exports/no_bid_docx", params={"version": 1}, headers=MANAGER).status_code == 404
    doc = Document(io.BytesIO(_download(c, tid, "no_bid_docx", 2)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Tender version: v2" in text and "NO_BID" in text and "REQ-PQ-06" in text
    assert "ISO 27001 expires before the deadline" in text  # the human decision comment
    assert "[1] " in text and "KB-CERT-002" in text  # numbered company evidence


def test_draft_docx_needs_approved_drafts(tender, tmp_path):
    c, tid = tender
    r = c.get(f"/tenders/{tid}/exports/draft_docx", params={"version": 2}, headers=MANAGER)
    assert r.status_code == 404  # drafting is Phase 10/11: nothing approved yet, so no file
    span = SourceSpan(doc_id="DOC-X", page=4, section_path="2.1 Eligible Bidders", text_span="single legal entity")
    draft = SectionDraft(draft_id="D1", tender_id=tid, tender_version=2, section_key="approach", title="Approach",
                         body_markdown="We bid as a single legal entity.", status=ApprovalStatus.approved,
                         claims=[Claim(claim_id="C1", text="single legal entity", citations=[span])])
    with doc_labels({"DOC-X": "main_rfp.pdf"}):
        path = write_draft_docx(tmp_path / "d.docx", tender_id=tid, version=2, drafts=[draft])
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "Tender version: v2" in text and "[1] main_rfp.pdf p.4, 2.1 Eligible Bidders" in text
